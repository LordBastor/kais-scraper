"""Convert ownership spreadsheets (ПИ / сгради / СОС) to TSV and combine them per category."""
from __future__ import annotations

import csv
import logging
import os
import shutil
from datetime import date, datetime
from pathlib import Path
from typing import Any

from .naming import OWNERSHIP_SUFFIXES, ownership_category
from .paths import Layout
from .util import run_parallel

log = logging.getLogger(__name__)


def _clean(value: Any) -> str:
    """Same as the old mlr step: newlines to spaces, collapse whitespace."""
    if value is None:
        return ""
    if isinstance(value, str):
        return " ".join(value.split())
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def convert_one(xlsx: Path, part: Path) -> int:
    """Write the first sheet of xlsx as a TSV (header included) to part; returns the row count."""
    from openpyxl import load_workbook

    part.parent.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    try:
        rows = wb.worksheets[0].iter_rows(values_only=True)
        first = next(rows, None)
        if first is None:
            return 0
        header = [_clean(v) for v in first]
        while header and header[-1] == "":
            header.pop()
        width = len(header)
        tmp = part.with_name(part.name + ".part")
        count = 0
        with tmp.open("w", encoding="utf-8", newline="") as fh:
            writer = csv.writer(fh, delimiter="\t", lineterminator="\n")
            writer.writerow(header)
            for row in rows:
                values = [_clean(v) for v in row[:width]]
                if not any(values):
                    continue
                values += [""] * (width - len(values))
                writer.writerow(values)
                count += 1
        os.replace(tmp, part)
        return count
    finally:
        wb.close()


def _combine(layout: Layout) -> None:
    for category in OWNERSHIP_SUFFIXES:
        parts = sorted((layout.ownership / "parts" / category).glob("*.tsv"))
        if not parts:
            continue
        out = layout.ownership / f"ownership-{category}.tsv"
        tmp = out.with_name(out.name + ".part")
        header: str | None = None
        mismatched: list[str] = []
        with tmp.open("w", encoding="utf-8", newline="") as dst:
            for part in parts:
                with part.open(encoding="utf-8", newline="") as src:
                    first = src.readline()
                    if header is None:
                        header = first
                        dst.write(first)
                    elif first != header:
                        mismatched.append(part.name)
                    shutil.copyfileobj(src, dst)
        os.replace(tmp, out)
        log.info("Wrote %s from %d files", out.name, len(parts))
        if mismatched:
            log.warning("%d %s files have a different header than the first one, e.g. %s",
                        len(mismatched), category, mismatched[:3])


def run_ownership(layout: Layout, workers: int) -> int:
    layout.ensure()
    jobs = []
    for xlsx in sorted(layout.extracted.rglob("*.xlsx")):
        category = ownership_category(xlsx.stem)
        if category is None:
            continue
        rel = xlsx.relative_to(layout.extracted).with_suffix(".tsv")
        part = layout.ownership / "parts" / category / "-".join(rel.parts)
        if part.exists() and part.stat().st_mtime >= xlsx.stat().st_mtime:
            continue
        jobs.append((xlsx, part))
    log.info("Converting %d changed spreadsheets", len(jobs))
    failures = run_parallel(convert_one, jobs, workers=workers, desc="xlsx", processes=True)
    _combine(layout)
    return len(failures)
